from rest_framework import views, permissions, status
from rest_framework.response import Response
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.contrib.auth import get_user_model
from django.http import HttpResponse

import csv
import io
from datetime import datetime, timedelta
from decimal import Decimal

from finance.models import Category, Income, Expense, Budget, SavingsGoal, Subscription, Notification
from finance.serializers import ExpenseSerializer, IncomeSerializer, NotificationSerializer

User = get_user_model()

class DashboardAnalyticsView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        user = request.user
        now = timezone.now()
        current_month = now.month
        current_year = now.year

        # Start of current week (assuming Monday is start of week)
        start_of_week = (now - timedelta(days=now.weekday())).date()
        start_of_month = now.replace(day=1).date()

        # Cumulative Metrics
        total_income = Income.objects.filter(user=user).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        total_expense = Expense.objects.filter(user=user).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        remaining_budget = total_income - total_expense
        total_savings = SavingsGoal.objects.filter(user=user).aggregate(total=Sum('current_amount'))['total'] or Decimal('0.00')

        # Spending limits
        current_month_spending = Expense.objects.filter(
            user=user, date__month=current_month, date__year=current_year
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        current_week_spending = Expense.objects.filter(
            user=user, date__gte=start_of_week
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        # 1. Monthly Expense Trend (Line Chart) - Last 6 months
        monthly_trend = []
        for i in range(5, -1, -1):
            date_check = now - timedelta(days=30 * i)
            m = date_check.month
            y = date_check.year
            month_spent = Expense.objects.filter(user=user, date__month=m, date__year=y).aggregate(total=Sum('amount'))['total'] or 0
            monthly_trend.append({
                "month": date_check.strftime("%b %Y"),
                "amount": float(month_spent)
            })

        # 2. Income vs Expense (Bar Chart) - Last 6 months
        income_vs_expense = []
        for i in range(5, -1, -1):
            date_check = now - timedelta(days=30 * i)
            m = date_check.month
            y = date_check.year
            m_spent = Expense.objects.filter(user=user, date__month=m, date__year=y).aggregate(total=Sum('amount'))['total'] or 0
            m_earned = Income.objects.filter(user=user, date__month=m, date__year=y).aggregate(total=Sum('amount'))['total'] or 0
            income_vs_expense.append({
                "month": date_check.strftime("%b %Y"),
                "income": float(m_earned),
                "expense": float(m_spent)
            })

        # 3. Expense Categories (Pie Chart) - Current Month
        category_spending = []
        categories = Category.objects.filter(Q(is_default=True) | Q(user=user))
        total_month_spent = float(current_month_spending) or 1.0 # avoid div by zero
        for cat in categories:
            cat_spent = Expense.objects.filter(
                user=user, category=cat, date__month=current_month, date__year=current_year
            ).aggregate(total=Sum('amount'))['total'] or 0
            if cat_spent > 0:
                category_spending.append({
                    "name": cat.name,
                    "value": float(cat_spent),
                    "percentage": round((float(cat_spent) / total_month_spent) * 100, 2)
                })

        # 4. Savings Trend (Area Chart) - Cumulative savings over last 6 months
        savings_trend = []
        running_savings = total_income - total_expense # Approximation
        for i in range(5, -1, -1):
            date_check = now - timedelta(days=30 * i)
            m = date_check.month
            y = date_check.year
            # Cumulative until this end of month
            limit_date = date_check.replace(day=28) # End of month approx
            inc_until = Income.objects.filter(user=user, date__lte=limit_date).aggregate(total=Sum('amount'))['total'] or 0
            exp_until = Expense.objects.filter(user=user, date__lte=limit_date).aggregate(total=Sum('amount'))['total'] or 0
            savings_trend.append({
                "month": date_check.strftime("%b %Y"),
                "savings": float(inc_until - exp_until)
            })

        # 5. Weekly Spending Analysis (current week daily breakdown)
        weekly_daily_breakdown = []
        for i in range(7):
            day_date = start_of_week + timedelta(days=i)
            day_spent = Expense.objects.filter(user=user, date=day_date).aggregate(total=Sum('amount'))['total'] or 0
            weekly_daily_breakdown.append({
                "day": day_date.strftime("%a"),
                "date": day_date.strftime("%Y-%m-%d"),
                "amount": float(day_spent)
            })

        # 6. Budget Utilization Progress
        budget_utilization = []
        monthly_budgets = Budget.objects.filter(user=user, month=current_month, year=current_year)
        for b in monthly_budgets:
            if b.budget_type == Budget.CATEGORY and b.category:
                spent = Expense.objects.filter(
                    user=user, category=b.category, date__month=current_month, date__year=current_year
                ).aggregate(total=Sum('amount'))['total'] or 0
                budget_utilization.append({
                    "name": b.category.name,
                    "budget": float(b.budget_amount),
                    "spent": float(spent),
                    "percentage": round((float(spent) / float(b.budget_amount)) * 100, 2) if b.budget_amount > 0 else 0
                })

        # Recent activities
        recent_expenses = Expense.objects.filter(user=user).order_by('-date', '-id')[:5]
        recent_incomes = Income.objects.filter(user=user).order_by('-date', '-id')[:5]
        recent_notifications = Notification.objects.filter(user=user).order_by('-created_at')[:5]

        return Response({
            "metrics": {
                "total_income": float(total_income),
                "total_expenses": float(total_expense),
                "remaining_budget": float(remaining_budget),
                "total_savings": float(total_savings),
                "current_month_spending": float(current_month_spending),
                "current_week_spending": float(current_week_spending),
            },
            "charts": {
                "monthly_expense_trend": monthly_trend,
                "income_vs_expense": income_vs_expense,
                "expense_categories": category_spending,
                "savings_trend": savings_trend,
                "weekly_spending_breakdown": weekly_daily_breakdown,
                "budget_utilization": budget_utilization
            },
            "recent_activity": {
                "expenses": ExpenseSerializer(recent_expenses, many=True).data,
                "incomes": IncomeSerializer(recent_incomes, many=True).data,
                "notifications": NotificationSerializer(recent_notifications, many=True).data
            }
        }, status=status.HTTP_200_OK)


class SmartInsightsView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        user = request.user
        now = timezone.now()
        current_month = now.month
        current_year = now.year

        # 1. Delta vs last month
        last_month = current_month - 1 if current_month > 1 else 12
        last_year = current_year if current_month > 1 else current_year - 1

        this_month_spent = Expense.objects.filter(
            user=user, date__month=current_month, date__year=current_year
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        last_month_spent = Expense.objects.filter(
            user=user, date__month=last_month, date__year=last_year
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        delta_percent = 0.0
        delta_type = "increased"
        if last_month_spent > 0:
            diff = this_month_spent - last_month_spent
            delta_percent = round(float((diff / last_month_spent) * 100), 2)
            if delta_percent < 0:
                delta_type = "decreased"
                delta_percent = abs(delta_percent)

        # 2. Highest spending category
        highest_cat = None
        highest_cat_spent = Decimal('0.00')
        categories = Category.objects.filter(Q(is_default=True) | Q(user=user))
        for cat in categories:
            cat_spent = Expense.objects.filter(
                user=user, category=cat, date__month=current_month, date__year=current_year
            ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
            if cat_spent > highest_cat_spent:
                highest_cat_spent = cat_spent
                highest_cat = cat.name

        # 3. Average daily expense
        days_in_month = now.day
        avg_daily = round(float(this_month_spent) / days_in_month, 2) if days_in_month > 0 else 0.0

        # 4. Savings rate
        this_month_income = Income.objects.filter(
            user=user, date__month=current_month, date__year=current_year
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        
        savings_rate = 0.0
        if this_month_income > 0:
            monthly_savings = this_month_income - this_month_spending
            savings_rate = round(float((monthly_savings / this_month_income) * 100), 2)

        # 5. Budget risks
        risks = []
        budgets = Budget.objects.filter(user=user, month=current_month, year=current_year, budget_type=Budget.CATEGORY)
        for b in budgets:
            actual = Expense.objects.filter(
                user=user, category=b.category, date__month=current_month, date__year=current_year
            ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
            if b.budget_amount > 0:
                pct = (actual / b.budget_amount) * 100
                if pct >= 80:
                    risks.append({
                        "category": b.category.name,
                        "budget": float(b.budget_amount),
                        "spent": float(actual),
                        "percentage": round(float(pct), 2),
                        "warning": f"Risk warning: You've utilized {round(float(pct), 1)}% of your '{b.category.name}' budget."
                    })

        return Response({
            "delta_comparison": {
                "percentage": delta_percent,
                "direction": delta_type,
                "text": f"Your spending has {delta_type} by {delta_percent}% compared to last month."
            },
            "highest_spending_category": {
                "category": highest_cat,
                "amount": float(highest_cat_spent)
            },
            "average_daily_expense": avg_daily,
            "savings_rate": savings_rate,
            "budget_risk_warnings": risks
        }, status=status.HTTP_200_OK)


class WeeklySummaryView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        user = request.user
        now = timezone.now()
        start_of_week = (now - timedelta(days=now.weekday())).date()

        # Weekly spending
        weekly_spending = Expense.objects.filter(
            user=user, date__gte=start_of_week
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

        # Weekly savings (approximation: weekly income - weekly expense)
        weekly_income = Income.objects.filter(
            user=user, date__gte=start_of_week
        ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        weekly_savings = weekly_income - weekly_spending

        # Top Categories this week
        top_categories = []
        categories = Category.objects.filter(Q(is_default=True) | Q(user=user))
        for cat in categories:
            spent = Expense.objects.filter(
                user=user, category=cat, date__gte=start_of_week
            ).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
            if spent > 0:
                top_categories.append({
                    "name": cat.name,
                    "amount": float(spent)
                })
        top_categories = sorted(top_categories, key=lambda x: x['amount'], reverse=True)[:3]

        # Daily spending breakdown
        daily_spending = []
        for i in range(7):
            d = start_of_week + timedelta(days=i)
            spent = Expense.objects.filter(user=user, date=d).aggregate(total=Sum('amount'))['total'] or 0
            daily_spending.append({
                "day": d.strftime("%A"),
                "date": d.strftime("%Y-%m-%d"),
                "amount": float(spent)
            })

        return Response({
            "weekly_spending": float(weekly_spending),
            "weekly_savings": float(weekly_savings),
            "top_categories": top_categories,
            "daily_spending_breakdown": daily_spending
        }, status=status.HTTP_200_OK)


class ReportExportView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        user = request.user
        export_format = request.query_params.get('format', 'csv').lower()
        report_range = request.query_params.get('range', 'monthly').lower()
        
        # Calculate dates
        now = timezone.now()
        if report_range == 'weekly':
            start_date = (now - timedelta(days=now.weekday())).date()
            end_date = now.date()
        elif report_range == 'yearly':
            start_date = now.replace(month=1, day=1).date()
            end_date = now.date()
        else: # monthly by default
            start_date = now.replace(day=1).date()
            end_date = now.date()

        # Handle custom overrides
        custom_start = request.query_params.get('start_date')
        custom_end = request.query_params.get('end_date')
        if custom_start:
            start_date = datetime.strptime(custom_start, "%Y-%m-%d").date()
        if custom_end:
            end_date = datetime.strptime(custom_end, "%Y-%m-%d").date()

        # Query data
        # For Admin exporting, if specified by a user query param, let admin see other data.
        target_user = user
        if user.role == 'ADMIN' and request.query_params.get('user_id'):
            try:
                target_user = User.objects.get(id=request.query_params.get('user_id'))
            except User.DoesNotExist:
                return Response({"error": "Target user not found"}, status=status.HTTP_404_NOT_FOUND)

        expenses = Expense.objects.filter(user=target_user, date__range=[start_date, end_date]).order_by('-date')
        incomes = Income.objects.filter(user=target_user, date__range=[start_date, end_date]).order_by('-date')

        # Compute totals
        tot_inc = incomes.aggregate(total=Sum('amount'))['total'] or 0
        tot_exp = expenses.aggregate(total=Sum('amount'))['total'] or 0
        net_savings = tot_inc - tot_exp

        # CSV format
        if export_format == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="BudgetFlow_Report_{report_range}_{start_date}_to_{end_date}.csv"'
            
            writer = csv.writer(response)
            writer.writerow(['BudgetFlow Report', f'Range: {start_date} to {end_date}', f'User: {target_user.username}'])
            writer.writerow([])
            writer.writerow(['INCOME ENTRIES'])
            writer.writerow(['Source', 'Amount', 'Date', 'Description'])
            for inc in incomes:
                writer.writerow([inc.source, inc.amount, inc.date, inc.description])
            writer.writerow(['Total Income', tot_inc])
            writer.writerow([])
            writer.writerow(['EXPENSE ENTRIES'])
            writer.writerow(['Name', 'Category', 'Amount', 'Payment Method', 'Date', 'Notes'])
            for exp in expenses:
                writer.writerow([exp.name, exp.category.name, exp.amount, exp.payment_method, exp.date, exp.notes])
            writer.writerow(['Total Expense', tot_exp])
            writer.writerow([])
            writer.writerow(['Net Savings', net_savings])
            return response

        # Excel format using openpyxl
        elif export_format == 'excel':
            import openpyxl
            wb = openpyxl.Workbook()
            ws_summary = wb.active
            ws_summary.title = "Summary"
            
            # Setup sheets
            ws_summary.append(['BudgetFlow Financial Summary'])
            ws_summary.append(['Period', f"{start_date} to {end_date}"])
            ws_summary.append(['Account Owner', target_user.username])
            ws_summary.append([])
            ws_summary.append(['Metric', 'Value'])
            ws_summary.append(['Total Income', float(tot_inc)])
            ws_summary.append(['Total Expense', float(tot_exp)])
            ws_summary.append(['Net Savings', float(net_savings)])

            ws_income = wb.create_sheet("Income")
            ws_income.append(['Source', 'Amount', 'Date', 'Description'])
            for inc in incomes:
                ws_income.append([inc.source, float(inc.amount), inc.date.strftime("%Y-%m-%d"), inc.description])

            ws_expense = wb.create_sheet("Expenses")
            ws_expense.append(['Name', 'Category', 'Amount', 'Payment Method', 'Date', 'Notes'])
            for exp in expenses:
                ws_expense.append([exp.name, exp.category.name, float(exp.amount), exp.payment_method, exp.date.strftime("%Y-%m-%d"), exp.notes])

            # Save in buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="BudgetFlow_Report_{report_range}_{start_date}_to_{end_date}.xlsx"'
            return response

        # PDF format using ReportLab
        elif export_format == 'pdf':
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
            elements = []

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'TitleStyle',
                parent=styles['Heading1'],
                fontSize=22,
                textColor=colors.HexColor('#1E293B'),
                spaceAfter=15
            )
            h2_style = ParagraphStyle(
                'H2Style',
                parent=styles['Heading2'],
                fontSize=14,
                textColor=colors.HexColor('#475569'),
                spaceBefore=15,
                spaceAfter=5
            )
            body_style = styles['Normal']

            elements.append(Paragraph("BudgetFlow Financial Report", title_style))
            elements.append(Paragraph(f"Period: {start_date} to {end_date}", body_style))
            elements.append(Paragraph(f"User Account: {target_user.name or target_user.username} ({target_user.email})", body_style))
            elements.append(Spacer(1, 15))

            # Summary Table
            summary_data = [
                ['Financial Metric', 'Amount'],
                ['Total Income', f"{target_user.currency} {tot_inc:.2f}"],
                ['Total Expenses', f"{target_user.currency} {tot_exp:.2f}"],
                ['Net Savings', f"{target_user.currency} {net_savings:.2f}"]
            ]
            t_summary = Table(summary_data, colWidths=[200, 150])
            t_summary.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#64748B')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 10),
                ('BOTTOMPADDING', (0,0), (-1,0), 8),
                ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#F8FAFC')),
                ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
            ]))
            elements.append(Paragraph("Overview Summary", h2_style))
            elements.append(t_summary)
            elements.append(Spacer(1, 15))

            # Income details
            elements.append(Paragraph("Income Entries", h2_style))
            income_table_data = [['Source', 'Amount', 'Date']]
            for inc in incomes[:10]: # Limit in PDF to avoid excessive pages
                income_table_data.append([inc.source, f"{target_user.currency} {inc.amount:.2f}", inc.date.strftime("%Y-%m-%d")])
            if len(incomes) > 10:
                income_table_data.append(['...', '...', '...'])
            
            t_inc = Table(income_table_data, colWidths=[150, 100, 100])
            t_inc.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#475569')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('FONTSIZE', (0,0), (-1,-1), 9),
            ]))
            elements.append(t_inc)
            elements.append(Spacer(1, 15))

            # Expenses details
            elements.append(Paragraph("Expense Entries (Top 15)", h2_style))
            expense_table_data = [['Name', 'Category', 'Amount', 'Date']]
            for exp in expenses[:15]:
                expense_table_data.append([exp.name[:20], exp.category.name, f"{target_user.currency} {exp.amount:.2f}", exp.date.strftime("%Y-%m-%d")])
            if len(expenses) > 15:
                expense_table_data.append(['...', '...', '...', '...'])
            
            t_exp = Table(expense_table_data, colWidths=[150, 100, 100, 100])
            t_exp.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#475569')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('FONTSIZE', (0,0), (-1,-1), 9),
            ]))
            elements.append(t_exp)

            doc.build(elements)
            buffer.seek(0)
            
            response = HttpResponse(buffer.read(), content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="BudgetFlow_Report_{report_range}_{start_date}_to_{end_date}.pdf"'
            return response

        return Response({"error": "Invalid format requested"}, status=status.HTTP_400_BAD_REQUEST)


class AdminAnalyticsView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        # Double check role
        if request.user.role != 'ADMIN':
            return Response({"error": "Admin access required"}, status=status.HTTP_403_FORBIDDEN)

        total_users = User.objects.filter(role=User.USER).count()
        active_users = User.objects.filter(is_active=True, role=User.USER).count()
        total_income_all = Income.objects.aggregate(total=Sum('amount'))['total'] or 0.0
        total_expense_all = Expense.objects.aggregate(total=Sum('amount'))['total'] or 0.0

        # System categories distribution
        category_stats = []
        categories = Category.objects.filter(is_default=True)
        for cat in categories:
            usage_count = Expense.objects.filter(category=cat).count()
            cat_sum = Expense.objects.filter(category=cat).aggregate(total=Sum('amount'))['total'] or 0.0
            category_stats.append({
                "category": cat.name,
                "transactions": usage_count,
                "amount": float(cat_sum)
            })

        # Monthly transaction velocity
        now = timezone.now()
        monthly_velocity = []
        for i in range(5, -1, -1):
            date_check = now - timedelta(days=30 * i)
            m = date_check.month
            y = date_check.year
            tx_count = Expense.objects.filter(date__month=m, date__year=y).count() + Income.objects.filter(date__month=m, date__year=y).count()
            monthly_velocity.append({
                "month": date_check.strftime("%b %Y"),
                "transactions": tx_count
            })

        return Response({
            "total_users": total_users,
            "active_users": active_users,
            "total_income": float(total_income_all),
            "total_expenses": float(total_expense_all),
            "category_analytics": category_stats,
            "monthly_velocity": monthly_velocity
        }, status=status.HTTP_200_OK)
